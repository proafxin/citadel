import logging
import re
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)

FENCE = "````"
GRID_HEADING = "## GRID"
META_HEADING = "## METADATA"
SPLIT = re.compile(r"(?<!\\)\|")
FORMULA_OPEN = "\u2039"
FORMULA_CLOSE = "\u203a"
FORMULA = re.compile(rf"^(.*?)\s*{FORMULA_OPEN}(.+){FORMULA_CLOSE}$", re.DOTALL)
MIN_FENCES = 2
META_LINE = re.compile(r"^- ([A-Z][A-Z ]*):\s*(.*)$")
EXTENT = re.compile(r"^- extent: `([A-Z]+)(\d+):([A-Z]+)(\d+)`")
SOURCE = re.compile(r"^- source: `(.+?)` sheet `(.+?)`")
WINDOW = re.compile(r"^- window: rows (\d+)-(\d+) of the sheet's rows (\d+)-(\d+)")
CELL_REF = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")
NAME_ENTRY = re.compile(r"^(\[[^\]]+\])?([^=]+)=(.*)$")


class Cell(BaseModel):
    row: int
    column: int
    ref: str
    value: str
    formula: str | None = None


class SheetMeta(BaseModel):
    merged: list[str] = Field(default_factory=list)
    types: str = ""
    formats: str = ""
    hidden_rows: list[int] = Field(default_factory=list)
    hidden_cols: list[str] = Field(default_factory=list)
    freeze: str | None = None
    autofilter: str | None = None
    list_objects: list[str] = Field(default_factory=list)
    defined_names: dict[str, str] = Field(default_factory=dict)
    validation: list[str] = Field(default_factory=list)
    cond_format: list[str] = Field(default_factory=list)
    comments: dict[str, str] = Field(default_factory=dict)
    hyperlinks: dict[str, str] = Field(default_factory=dict)
    bold: str = ""
    blank_cols: list[str] = Field(default_factory=list)


class SheetDump(BaseModel):
    workbook: str
    sheet: str
    first_row: int
    last_row: int
    first_col: int
    last_col: int
    sheet_first_row: int = 0
    sheet_last_row: int = 0
    columns: list[int] = Field(default_factory=list)
    cells: dict[int, dict[int, Cell]] = Field(default_factory=dict)
    lines: dict[int, str] = Field(default_factory=dict)
    meta: SheetMeta = Field(default_factory=SheetMeta)

    @property
    def windowed(self) -> bool:
        return (self.sheet_first_row, self.sheet_last_row) != (self.first_row, self.last_row)

    @property
    def extent(self) -> str:
        if not self.first_col or not self.last_col:
            return f"rows {self.first_row}-{self.last_row}"
        return f"{get_column_letter(self.first_col)}{self.first_row}:{get_column_letter(self.last_col)}{self.last_row}"

    def cell(self, row: int, column: int) -> Cell | None:
        return self.cells.get(row, {}).get(column)


def unescape(text: str) -> str:
    return text.replace("\\|", "|")


def split_ruler(line: str) -> list[int]:
    parts = SPLIT.split(line)[1:]
    columns = []
    for part in parts:
        letters = "".join(ch for ch in part if ch.isalpha())
        if letters:
            columns.append(column_index_from_string(letters))
    return columns


def parse_row(line: str, columns: list[int]) -> tuple[int, dict[int, Cell]]:
    parts = SPLIT.split(line)
    row = int(parts[0])
    out: dict[int, Cell] = {}
    for offset, raw in enumerate(parts[1:]):
        if offset >= len(columns):
            break
        text = unescape(raw)
        if not text:
            continue
        column = columns[offset]
        match = FORMULA.match(text)
        value = match.group(1) if match else text
        formula = match.group(2) if match else None
        out[column] = Cell(
            row=row, column=column, ref=f"{get_column_letter(column)}{row}", value=value, formula=formula
        )
    return row, out


def parse_defined_names(payload: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for entry in payload.split(" | "):
        match = NAME_ENTRY.match(entry.strip())
        if match:
            scope = match.group(1) or ""
            out[f"{scope}{match.group(2).strip()}"] = match.group(3).strip()
    return out


def parse_pairs(payload: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for entry in payload.split(" | "):
        key, sep, value = entry.strip().partition("=")
        if sep:
            out[key] = value
    return out


def parse_hidden(payload: str) -> tuple[list[int], list[str]]:
    rows_part, _, cols_part = payload.partition("cols")
    rows = [int(n) for n in re.findall(r"\d+", rows_part.replace("rows", ""))]
    cols = re.findall(r"[A-Z]+", cols_part)
    return rows, cols


def blank(payload: str) -> bool:
    return payload.strip() in {"", "-", "[]", "None"}


def parse_list(payload: str) -> list[str]:
    return [part.strip() for part in payload.split(" | ")]


def parse_list_objects(payload: str) -> list[str]:
    return re.findall(r"'([^']+)'", payload)


def parse_cols(payload: str) -> list[str]:
    return re.findall(r"[A-Z]+", payload)


RAW_FIELDS = {"TYPES": "types", "FORMATS": "formats", "BOLD": "bold"}
SCALAR_FIELDS = {"FREEZE": "freeze", "AUTOFILTER": "autofilter"}
PARSED_FIELDS = {
    "MERGED": ("merged", parse_list),
    "VALIDATION": ("validation", parse_list),
    "CONDFMT": ("cond_format", parse_list),
    "LISTOBJECTS": ("list_objects", parse_list_objects),
    "DEFINED NAMES": ("defined_names", parse_defined_names),
    "COMMENTS": ("comments", parse_pairs),
    "HYPERLINKS": ("hyperlinks", parse_pairs),
    "BLANK COLS": ("blank_cols", parse_cols),
}


def apply_meta(meta: SheetMeta, key: str, payload: str) -> None:
    if key in RAW_FIELDS:
        setattr(meta, RAW_FIELDS[key], payload)
        return
    if key == "HIDDEN":
        meta.hidden_rows, meta.hidden_cols = parse_hidden(payload)
        return
    if blank(payload):
        return
    if key in SCALAR_FIELDS:
        setattr(meta, SCALAR_FIELDS[key], payload)
        return
    if key in PARSED_FIELDS:
        field, parser = PARSED_FIELDS[key]
        setattr(meta, field, parser(payload))


def parse_meta(lines: list[str]) -> SheetMeta:
    meta = SheetMeta()
    for line in lines:
        match = META_LINE.match(line)
        if match:
            apply_meta(meta, match.group(1).strip(), match.group(2).strip())
    return meta


def parse_header(lines: list[str]) -> tuple[str, str, tuple[int, int, int, int], tuple[int, int]]:
    workbook = ""
    sheet = ""
    box = (0, 0, 0, 0)
    sheet_rows = (0, 0)
    for line in lines:
        window = WINDOW.match(line)
        if window:
            sheet_rows = (int(window.group(3)), int(window.group(4)))
        source = SOURCE.match(line)
        if source:
            workbook, sheet = source.group(1), source.group(2)
        extent = EXTENT.match(line)
        if extent:
            box = (
                int(extent.group(2)),
                int(extent.group(4)),
                column_index_from_string(extent.group(1)),
                column_index_from_string(extent.group(3)),
            )
    return workbook, sheet, box, sheet_rows


def parse_dump(text: str) -> SheetDump:
    lines = text.splitlines()
    workbook, sheet, box, sheet_rows = parse_header(lines)
    grid_start = next((i for i, line in enumerate(lines) if line.strip() == GRID_HEADING), -1)
    if not sheet:
        logger.warning("dump has no source header, sheet identity unknown")
    meta_start = next((i for i, line in enumerate(lines) if line.strip() == META_HEADING), len(lines))
    dump = SheetDump(
        workbook=workbook,
        sheet=sheet,
        first_row=box[0],
        last_row=box[1],
        first_col=box[2],
        last_col=box[3],
        sheet_first_row=sheet_rows[0] or box[0],
        sheet_last_row=sheet_rows[1] or box[1],
        meta=parse_meta(lines[meta_start:]),
    )
    if grid_start < 0:
        logger.warning("no GRID section in dump for %s/%s", workbook, sheet)
        return dump
    fences = [i for i, line in enumerate(lines[grid_start:meta_start], grid_start) if line.strip() == FENCE]
    if len(fences) < MIN_FENCES:
        logger.warning("unterminated grid fence in dump for %s/%s", workbook, sheet)
        return dump
    body = lines[fences[0] + 1 : fences[1]]
    if not body:
        return dump
    dump.columns = split_ruler(body[0])
    for line in body[1:]:
        if not line or not line.split("|")[0].strip().isdigit():
            continue
        row, cells = parse_row(line, dump.columns)
        dump.lines[row] = line
        if cells:
            dump.cells[row] = cells
    return dump


def load_dump(path: Path) -> SheetDump:
    return parse_dump(path.read_text(encoding="utf-8"))
