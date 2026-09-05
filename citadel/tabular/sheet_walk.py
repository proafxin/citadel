import logging
from functools import lru_cache
from pathlib import Path

from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from citadel.llm import call_structured
from citadel.tabular.sheet_dump import SheetDump
from citadel.tabular.sheet_structure import (
    MetadataItem,
    MetadataKind,
    Orientation,
    Region,
    RowRole,
    RowSpan,
    SheetStructure,
)

logger = logging.getLogger(__name__)

CONTEXT_BEFORE = 8
SCAN_ROWS = 25
CONTEXT_AFTER = 12
WALK_MAX_TOKENS = 4096
MAX_SWEEPS = 4
SCAN_PROMPT = Path(__file__).resolve().parents[2] / "prompts" / "sheet_scan.md"
ROW_PROMPT = Path(__file__).resolve().parents[2] / "prompts" / "sheet_row.md"
REVIEW_PROMPT = Path(__file__).resolve().parents[2] / "prompts" / "sheet_review.md"


@lru_cache(maxsize=1)
def scan_prompt() -> str:
    return SCAN_PROMPT.read_text(encoding="utf-8")


@lru_cache(maxsize=1)
def row_prompt() -> str:
    return ROW_PROMPT.read_text(encoding="utf-8")


@lru_cache(maxsize=1)
def review_prompt() -> str:
    return REVIEW_PROMPT.read_text(encoding="utf-8")


class ScanBlock(BaseModel):
    first_row: int
    last_row: int
    kind: MetadataKind
    summary: str = ""


class ScanTable(BaseModel):
    first_row: int
    header_rows: list[int] = Field(default_factory=list)
    first_col: int = 0
    last_col: int = 0
    orientation: Orientation = Orientation.ROW_RECORDS


class ScanVerdict(BaseModel):
    tables: list[ScanTable] = Field(default_factory=list)
    blocks: list[ScanBlock] = Field(default_factory=list)


class RowVerdict(BaseModel):
    role: RowRole
    continues: bool
    summary: str = ""


class OpenTable(BaseModel):
    first_row: int
    last_row: int
    first_col: int
    last_col: int
    orientation: Orientation = Orientation.ROW_RECORDS
    header_rows: list[int] = Field(default_factory=list)
    shapes: dict[str, RowRole] = Field(default_factory=dict)
    roles: dict[int, RowRole] = Field(default_factory=dict)


class WalkBlock(BaseModel):
    first_row: int
    last_row: int
    kind: MetadataKind
    summary: str


class WalkResult(BaseModel):
    tables: list[OpenTable] = Field(default_factory=list)
    blocks: list[WalkBlock] = Field(default_factory=list)
    calls: int = 0
    extended: int = 0
    unclaimed: list[str] = Field(default_factory=list)


def overlaps(a: OpenTable, b: OpenTable) -> bool:
    shared = set(a.roles) & set(b.roles)
    return bool(shared) and a.first_col <= b.last_col and b.first_col <= a.last_col


def merge_tables(tables: list[OpenTable]) -> list[OpenTable]:
    merged: list[OpenTable] = []
    for table in sorted(tables, key=lambda t: (t.first_row, t.first_col)):
        for other in merged:
            if not overlaps(table, other):
                continue
            other.first_row = min(other.first_row, table.first_row)
            other.last_row = max(other.last_row, table.last_row)
            other.first_col = min(other.first_col, table.first_col)
            other.last_col = max(other.last_col, table.last_col)
            other.header_rows = sorted(set(other.header_rows) | set(table.header_rows))
            other.shapes.update(table.shapes)
            for row, role in table.roles.items():
                other.roles.setdefault(row, role)
            break
        else:
            merged.append(table)
    return merged


def has_records(table: OpenTable) -> bool:
    return any(role == RowRole.BODY for role in table.roles.values())


def claimed_cells(result: WalkResult, dump: SheetDump) -> set[tuple[int, int]]:
    held: set[tuple[int, int]] = set()
    for table in result.tables:
        for row in table.roles:
            held.update(
                (row, col)
                for col in range(table.first_col, table.last_col + 1)
                if dump.cell(row, col) is not None
            )
    for block in result.blocks:
        for row in range(block.first_row, block.last_row + 1):
            held.update((row, col) for col in dump.cells.get(row, {}))
    return held


def unclaimed_rows(result: WalkResult, dump: SheetDump) -> list[int]:
    held = claimed_cells(result, dump)
    return sorted({row for row, columns in dump.cells.items() for col in columns if (row, col) not in held})


def cell_kind(value: str) -> str:
    text = value.strip()
    if text.startswith("#") and text.endswith(("!", "?")):
        return "n"
    bare = text.replace(",", "").replace("$", "").replace("%", "").strip()
    if bare.startswith("(") and bare.endswith(")"):
        bare = bare[1:-1]
    try:
        float(bare)
    except ValueError:
        return "t"
    return "n"


def row_shape(dump: SheetDump, row: int, first_col: int, last_col: int) -> str:
    parts = []
    for col in range(first_col, last_col + 1):
        cell = dump.cell(row, col)
        if cell is not None:
            parts.append(f"{get_column_letter(col)}:{cell_kind(cell.value)}")
    return ",".join(parts)


def slice_lines(dump: SheetDump, first_row: int, last_row: int) -> str:
    rows = [row for row in range(first_row, last_row + 1) if row in dump.lines]
    counts = "  ".join(f"{row}:{len(dump.cells.get(row, {}))}" for row in rows)
    body = "\n".join(dump.lines[row] for row in rows)
    return f"cells per row: {counts}\n\n{body}"


def context(dump: SheetDump, row: int) -> str:
    first = max(dump.first_row, row - CONTEXT_BEFORE)
    last = min(dump.last_row, row + CONTEXT_AFTER)
    return slice_lines(dump, first, last)


def describe(table: OpenTable) -> str:
    shapes = "; ".join(f"{shape or '(empty)'} -> {role.value}" for shape, role in table.shapes.items())
    return (
        f"An open table runs from row {table.first_row}, columns "
        f"{get_column_letter(table.first_col)}-{get_column_letter(table.last_col)}, "
        f"header rows {table.header_rows or 'none'}, last row settled so far {table.last_row}.\n"
        f"Row shapes already settled in it: {shapes or 'none'}."
    )


async def scan_for_table(dump: SheetDump, first_row: int, last_row: int) -> ScanVerdict:
    prompt = "\n\n".join(
        [
            scan_prompt(),
            f"## Worksheet {dump.workbook} / {dump.sheet}, rows {first_row}-{last_row}",
            slice_lines(dump, first_row, last_row),
            f"Looking only at rows {first_row} to {last_row}, give every table that starts in them: the row it "
            "starts at, the rows that name its columns, and the columns it occupies. Tables that sit side by "
            "side start at the same row in different columns; report each one. Describe everything that is not "
            "a table as blocks with their kind and a one-line summary.",
        ]
    )
    raw = await call_structured(prompt, ScanVerdict.model_json_schema(), max_tokens=WALK_MAX_TOKENS)
    return ScanVerdict.model_validate(raw)


async def ask_row(dump: SheetDump, row: int, table: OpenTable) -> RowVerdict:
    prompt = "\n\n".join(
        [
            row_prompt(),
            f"## Worksheet {dump.workbook} / {dump.sheet}, rows around {row}",
            context(dump, row),
            describe(table),
            f"Row {row} holds cells in columns {row_shape(dump, row, table.first_col, table.last_col) or 'none'}, "
            "a pattern not settled in this table yet. Give its role, and say whether the table continues "
            "through it or has already ended above it.",
        ]
    )
    raw = await call_structured(prompt, RowVerdict.model_json_schema(), max_tokens=WALK_MAX_TOKENS)
    return RowVerdict.model_validate(raw)


class ReviewVerdict(BaseModel):
    tables: list[ScanTable] = Field(default_factory=list)


async def review_blocks(dump: SheetDump, result: WalkResult) -> ReviewVerdict:
    listing = "\n".join(
        f"rows {b.first_row}-{b.last_row} {b.kind.value}: {b.summary}"
        for b in sorted(result.blocks, key=lambda x: x.first_row)
    )
    rows = sorted({r for b in result.blocks for r in range(b.first_row, b.last_row + 1) if r in dump.lines})
    prompt = "\n\n".join(
        [
            review_prompt(),
            f"## Worksheet {dump.workbook} / {dump.sheet}",
            "## Rows held by no table\n\n" + "\n".join(dump.lines[r] for r in rows),
            "## How they were described\n\n" + listing,
            "Give every table among these rows that was described as something else.",
        ]
    )
    raw = await call_structured(prompt, ReviewVerdict.model_json_schema(), max_tokens=WALK_MAX_TOKENS)
    return ReviewVerdict.model_validate(raw)


def open_from(scan: ScanTable, dump: SheetDump, settled: list[OpenTable] | None = None) -> OpenTable:
    headers = [h for h in scan.header_rows if h >= scan.first_row] or [scan.first_row]
    table = OpenTable(
        first_row=scan.first_row,
        last_row=max(headers),
        first_col=scan.first_col or dump.first_col,
        last_col=scan.last_col or dump.last_col,
        header_rows=headers,
        orientation=scan.orientation,
    )
    for header in headers:
        table.roles[header] = RowRole.HEADER
    for known in settled or []:
        if table.first_col <= known.last_col and known.first_col <= table.last_col:
            table.shapes.update(known.shapes)
    return table


async def sweep(
    dump: SheetDump,
    result: WalkResult,
    rows: list[int],
    max_calls: int,
    seed: list[ScanTable] | None = None,
) -> None:
    tables: list[OpenTable] = [open_from(scan, dump, result.tables) for scan in seed or []]
    pending = [r for r in rows if not tables or r > max(t.last_row for t in tables)]
    while pending and result.calls < max_calls:
        row = pending.pop(0)
        if not tables:
            window = [row, *[r for r in pending if r <= row + SCAN_ROWS - 1]]
            verdict = await scan_for_table(dump, row, window[-1])
            result.calls += 1
            result.blocks.extend(
                WalkBlock(first_row=b.first_row, last_row=b.last_row, kind=b.kind, summary=b.summary)
                for b in verdict.blocks
            )
            found = [t for t in verdict.tables if row <= t.first_row <= window[-1]]
            if not found:
                pending = [r for r in pending if r > window[-1]]
                continue
            tables.extend(open_from(start, dump, result.tables) for start in found)
            top = max(t.last_row for t in tables)
            pending = [r for r in pending if r > top]
            continue
        for table in list(tables):
            if row <= table.last_row:
                continue
            shape = row_shape(dump, row, table.first_col, table.last_col)
            if not shape:
                table.roles[row] = RowRole.BLANK
                continue
            known = table.shapes.get(shape)
            if known is not None:
                table.roles[row] = known
                table.last_row = row
                result.extended += 1
                continue
            if result.calls >= max_calls:
                break
            verdict = await ask_row(dump, row, table)
            result.calls += 1
            if not verdict.continues:
                if has_records(table):
                    result.tables.append(table)
                tables.remove(table)
                pending.insert(0, row)
                break
            table.shapes[shape] = verdict.role
            table.roles[row] = verdict.role
            table.last_row = row
    result.tables.extend(table for table in tables if has_records(table))


async def walk_sheet(dump: SheetDump, max_calls: int = 400) -> WalkResult:
    result = WalkResult()
    rows = sorted(dump.cells)
    for index in range(MAX_SWEEPS):
        if not rows or result.calls >= max_calls:
            break
        before = len(rows)
        await sweep(dump, result, rows, max_calls)
        result.tables = merge_tables(result.tables)
        rows = unclaimed_rows(result, dump)
        logger.info(
            "%s/%s sweep %d: %d tables, %d blocks, %d calls, %d rows still unclaimed",
            dump.workbook,
            dump.sheet,
            index + 1,
            len(result.tables),
            len(result.blocks),
            result.calls,
            len(rows),
        )
        if len(rows) >= before:
            break
    if result.blocks and result.calls < max_calls:
        verdict = await review_blocks(dump, result)
        result.calls += 1
        promoted = [t for t in verdict.tables if t.first_row in dump.cells]
        if promoted:
            claimed = {
                row
                for t in promoted
                for row in range(t.first_row, dump.last_row + 1)
                if row in dump.cells
            }
            result.blocks = [b for b in result.blocks if b.first_row not in claimed]
            logger.info(
                "%s/%s review promoted %d block(s) to tables", dump.workbook, dump.sheet, len(promoted)
            )
            await sweep(dump, result, sorted(claimed), max_calls, seed=promoted)
            result.tables = merge_tables(result.tables)
    held = claimed_cells(result, dump)
    result.unclaimed = [
        f"{get_column_letter(col)}{row}"
        for row, columns in sorted(dump.cells.items())
        for col in sorted(columns)
        if (row, col) not in held
    ]
    logger.info(
        "walked %s/%s: %d tables, %d blocks, %d calls, %d rows extended free, %d cells unclaimed",
        dump.workbook,
        dump.sheet,
        len(result.tables),
        len(result.blocks),
        result.calls,
        result.extended,
        len(result.unclaimed),
    )
    return result


def role_spans(table: OpenTable) -> list[RowSpan]:
    spans: list[RowSpan] = []
    for row in sorted(table.roles):
        role = table.roles[row]
        if spans and spans[-1].role == role and spans[-1].to_row == row - 1:
            spans[-1].to_row = row
        else:
            spans.append(RowSpan(from_row=row, to_row=row, role=role))
    return spans


def walk_structure(dump: SheetDump, result: WalkResult) -> SheetStructure:
    regions = [
        Region(
            region_id=f"r{index + 1}",
            first_row=min(table.roles) if table.roles else table.first_row,
            last_row=max(table.roles) if table.roles else table.last_row,
            first_col=table.first_col,
            last_col=table.last_col,
            row_spans=role_spans(table),
            orientation=table.orientation,
        )
        for index, table in enumerate(sorted(result.tables, key=lambda t: (t.first_row, t.first_col)))
        if table.roles
    ]
    metadata = [
        MetadataItem(
            kind=block.kind,
            first_row=block.first_row,
            last_row=block.last_row,
            first_col=dump.first_col,
            last_col=dump.last_col,
            summary=block.summary,
        )
        for block in result.blocks
    ]
    return SheetStructure(
        workbook=dump.workbook,
        sheet=dump.sheet,
        regions=regions,
        metadata=metadata,
        unresolved=result.unclaimed[:5],
        rounds=result.calls,
    )
