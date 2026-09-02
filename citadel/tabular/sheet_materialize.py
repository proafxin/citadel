import logging
import time

from openpyxl.utils import get_column_letter

from citadel.schemas.table import TableStructure as GridStructure
from citadel.services.excel import (
    DUMP_VERSION,
    Cell,
    SheetExtraction,
    SheetItem,
    SheetText,
    cell_value,
    render_sheet_dump,
)
from citadel.tabular.materialize import MaterializedTable, materialize
from citadel.tabular.sheet_agent import extract_sheet, fits_context
from citadel.tabular.sheet_dump import SheetDump, parse_dump
from citadel.tabular.sheet_structure import (
    Orientation,
    Region,
    RowRole,
    SheetStructure,
    body_row_numbers,
    letter_index,
    rows_by_role,
)

logger = logging.getLogger(__name__)

DEFAULT_BAND_NAME = "section"


def cell_index(cells: list[Cell]) -> dict[tuple[int, int], Cell]:
    return {(cell.row, cell.col): cell for cell in cells}


def render(cell: Cell | None) -> str:
    return "" if cell is None or (value := cell_value(cell)) is None else str(value)


def region_rows(region: Region) -> list[int]:
    kept = {RowRole.HEADER, RowRole.BODY, RowRole.TOTALS, RowRole.BAND_LABEL}
    rows = {row for role in kept for row in rows_by_role(region, role)}
    return [row for row in sorted(rows) if region.first_row <= row <= region.last_row]


def region_columns(index: dict[tuple[int, int], Cell], region: Region, rows: list[int]) -> list[int]:
    return [
        col
        for col in range(region.first_col, region.last_col + 1)
        if any((row, col) in index for row in rows)
    ]


def build_grid(index: dict[tuple[int, int], Cell], rows: list[int], columns: list[int]) -> list[list[str]]:
    return [[render(index.get((row, col))) for col in columns] for row in rows]


def band_target(region: Region, columns: list[int]) -> int | None:
    wanted = letter_index(region.band_column)
    return columns.index(wanted) if wanted in columns else None


def band_for_rows(region: Region) -> dict[int, str]:
    return {
        row: band.label
        for band in region.bands
        for row in range(min(band.from_row, band.to_row), max(band.from_row, band.to_row) + 1)
    }


def grid_structure(region: Region, rows: list[int], width: int) -> GridStructure | None:
    positions = {row: offset for offset, row in enumerate(rows)}
    body = [positions[row] for row in body_row_numbers(region) if row in positions]
    if not body:
        return None
    header = [positions[row] for row in rows_by_role(region, RowRole.HEADER) if row in positions]
    extra = [
        positions[row]
        for role in (RowRole.TOTALS, RowRole.BAND_LABEL)
        for row in rows_by_role(region, role)
        if row in positions
    ]
    return GridStructure(
        transposed=region.orientation == Orientation.MATRIX,
        col_start=0,
        col_end=width - 1,
        header_rows=header or None,
        data_start=min(body),
        data_end=max(body),
        metadata_rows=sorted(extra) or None,
        columns=[column.name for column in region.columns] or None,
        notes=region.support or None,
    )


def anchors(region: Region, columns: list[int], structure: SheetStructure) -> dict:
    return {
        "region_id": region.region_id,
        "min_row": region.first_row,
        "min_col": region.first_col,
        "max_row": region.last_row,
        "max_col": region.last_col,
        "columns": columns,
        "orientation": region.orientation.value,
        "confidence": region.confidence.value,
        "support": region.support,
        "key_column": region.key_column,
        "sheet_rows": {
            role.value: rows_by_role(region, role)
            for role in (RowRole.HEADER, RowRole.BODY, RowRole.TOTALS, RowRole.BAND_LABEL)
        },
        "group_name": region.group_name,
        "band_name": region.band_name,
        "band_column": region.band_column,
        "bands": [[band.label, band.from_row, band.to_row] for band in region.bands],
        "metadata": [
            {"kind": item.kind.value, "summary": item.summary, "rows": [item.first_row, item.last_row]}
            for item in structure.metadata
            if region.region_id in item.region_ids
        ],
        "rounds": structure.rounds,
        "unresolved": structure.unresolved,
        "dump_version": DUMP_VERSION,
    }


def apply_names(built: MaterializedTable, region: Region, columns: list[int]) -> None:
    by_letter = {column.letter.strip().upper(): column.name for column in region.columns if column.name}
    by_index = {column.index: column.name for column in region.columns if column.name}
    for offset, column in enumerate(built.columns[: len(columns)]):
        col = columns[offset]
        name = by_index.get(col) or by_letter.get(get_column_letter(col))
        if name:
            column.header = name


def apply_groups(built: MaterializedTable, region: Region, columns: list[int]) -> None:
    by_index = {column.index: column.group for column in region.columns if column.group}
    by_letter = {column.letter.strip().upper(): column.group for column in region.columns if column.group}
    for offset, column in enumerate(built.columns[: len(columns)]):
        col = columns[offset]
        column.group = by_index.get(col) or by_letter.get(get_column_letter(col))


def region_title(region: Region, structure: SheetStructure) -> str | None:
    titles = [
        item.summary
        for item in structure.metadata
        if region.region_id in item.region_ids and item.kind.value in {"title", "caption"}
    ]
    return titles[0] if titles else None


def materialize_region(
    sheet: SheetExtraction,
    index: dict[tuple[int, int], Cell],
    region: Region,
    structure: SheetStructure,
) -> MaterializedTable | None:
    rows = region_rows(region)
    if not rows:
        logger.warning("skipping %s: no rows inside its own extent", region.region_id)
        return None
    columns = region_columns(index, region, rows)
    if not columns:
        logger.warning("skipping %s: no populated columns", region.region_id)
        return None
    grid = build_grid(index, rows, columns)
    bands = band_for_rows(region)
    target = band_target(region, columns)
    if bands:
        body = set(body_row_numbers(region))
        for offset, row in enumerate(rows):
            value = bands.get(row, "") if row in body else ""
            if target is None:
                grid[offset] = [*grid[offset], value]
            elif value:
                grid[offset][target] = value
    width = len(columns) + (1 if bands and target is None else 0)
    structure_spec = grid_structure(region, rows, width)
    if structure_spec is None:
        logger.warning("skipping %s: no rows holding records", region.region_id)
        return None
    formulas = sorted(
        {
            cell.formula
            for row in rows
            for col in columns
            if (cell := index.get((row, col))) is not None and cell.formula is not None
        }
    )
    structure_spec.title = region_title(region, structure)
    structure_spec.caption = structure_spec.title
    built = materialize(
        grid,
        structure_spec,
        sheet_no=sheet.sheet_no,
        formulas=formulas or None,
        extra_notes=[item.summary for item in structure.metadata] or None,
        anchors=anchors(region, columns, structure),
    )
    apply_names(built, region, columns)
    apply_groups(built, region, columns)
    if bands and built.columns:
        slot = built.columns[target if target is not None else -1]
        slot.header = region.band_name or DEFAULT_BAND_NAME
        slot.group = None
    return built


def materialize_sheet(
    sheet: SheetExtraction, dump: SheetDump, structure: SheetStructure
) -> list[tuple[int, MaterializedTable]]:
    index = cell_index(sheet.cells)
    produced: list[tuple[int, MaterializedTable]] = []
    for region in structure.regions:
        built = materialize_region(sheet, index, region, structure)
        if built is not None:
            produced.append((len(produced) + 1, built))
    return produced


def unattached_notes(structure: SheetStructure) -> str:
    return "\n".join(
        f"rows {item.first_row}-{item.last_row} {item.kind.value}: {item.summary}"
        for item in structure.metadata
        if not item.region_ids
    )


async def structure_sheet(sheet: SheetExtraction, workbook: str) -> list[tuple[int, SheetItem]]:
    queued = time.time()
    text = render_sheet_dump(sheet, workbook)
    dump = parse_dump(text)
    if not dump.last_row:
        logger.info("%s/%s has no populated cells, nothing to structure", workbook, sheet.sheet_name)
        return []
    logger.info(
        "structuring %s/%s extent=%s cells=%d dump_chars=%d",
        workbook,
        sheet.sheet_name,
        dump.extent,
        len(sheet.cells),
        len(text),
    )
    fits, needed, allowed = fits_context(dump, text)
    if not fits:
        logger.warning(
            "skipping oversized sheet %s/%s: %d input tokens exceeds %d",
            workbook,
            sheet.sheet_name,
            needed,
            allowed,
        )
        return []
    started = time.time()
    structure = await extract_sheet(dump, text)
    elapsed = time.time() - started
    materialized = materialize_sheet(sheet, dump, structure)
    items: list[tuple[int, SheetItem]] = list(materialized)
    notes = unattached_notes(structure)
    if notes:
        items.append((len(items) + 1, SheetText(sheet_no=sheet.sheet_no, text=notes)))
    logger.info(
        "structured %s/%s in %.1fs (queued %.1fs): %d regions reported, %d materialized, %d rows",
        workbook,
        sheet.sheet_name,
        elapsed,
        started - queued,
        len(structure.regions),
        len(materialized),
        sum(table.n_rows for _, table in materialized),
    )
    if structure.unresolved:
        logger.warning("%s/%s left unresolved: %s", workbook, sheet.sheet_name, ", ".join(structure.unresolved[:5]))
    return items
