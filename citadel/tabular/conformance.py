from collections.abc import Sequence
from enum import StrEnum

from openpyxl.utils import get_column_letter
from pydantic import BaseModel

type DTypeGrid = Sequence[Sequence[str | None]]

ERROR = "e"
STRING = "s"
NUMBER = "n"


def normalise(dtype: str | None, value: object) -> str | None:
    if dtype != STRING:
        return dtype
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        float(text)
    except ValueError:
        return STRING
    return NUMBER


class RowState(StrEnum):
    MATCH = "match"
    BLANK = "blank"
    DIFFER = "differ"


class Run(BaseModel):
    first_row: int
    last_row: int
    state: RowState
    columns: list[str]
    fewest: int
    most: int
    deviations: list[str]
    first_value: str


class ConformanceMap(BaseModel):
    reference: dict[str, str]
    runs: list[Run]


def _dtype(grid: DTypeGrid, row: int, col: int) -> str | None:
    if not 0 <= row < len(grid):
        return None
    line = grid[row]
    return line[col] if 0 <= col < len(line) else None


def column_reference(grid: DTypeGrid, data_start: int, first_col: int, last_col: int) -> dict[int, str]:
    reference: dict[int, str] = {}
    for col in range(first_col, last_col + 1):
        for row in range(data_start, len(grid)):
            dtype = _dtype(grid, row, col)
            if dtype is not None and dtype != ERROR:
                reference[col] = dtype
                break
    return reference


def last_populated(grid: DTypeGrid, data_start: int, first_col: int, last_col: int) -> int:
    last = data_start - 1
    for row in range(data_start, len(grid)):
        if any(_dtype(grid, row, col) is not None for col in range(first_col, last_col + 1)):
            last = row
    return last


def row_state(
    reference: dict[int, str], grid: DTypeGrid, row: int, first_col: int, last_col: int
) -> tuple[RowState, list[str], list[str]]:
    populated: list[str] = []
    deviations: list[str] = []
    for col in range(first_col, last_col + 1):
        dtype = _dtype(grid, row, col)
        if dtype is None:
            continue
        populated.append(get_column_letter(col + 1))
        if dtype != ERROR and col in reference and dtype != reference[col]:
            deviations.append(f"{get_column_letter(col + 1)}={dtype}")
    if not populated:
        return RowState.BLANK, [], []
    return (RowState.DIFFER if deviations else RowState.MATCH), deviations, populated


def conformance_map(
    grid: DTypeGrid, data_start: int, first_col: int, last_col: int, labels: Sequence[str] | None = None
) -> ConformanceMap:
    reference = column_reference(grid, data_start, first_col, last_col)
    end = last_populated(grid, data_start, first_col, last_col)
    runs: list[Run] = []
    row = data_start
    while row <= end:
        state, deviations, populated = row_state(reference, grid, row, first_col, last_col)
        start = row
        seen = set(populated)
        fewest = most = len(populated)
        while row + 1 <= end:
            nxt_state, nxt_dev, nxt_pop = row_state(reference, grid, row + 1, first_col, last_col)
            if (nxt_state, nxt_dev) != (state, deviations):
                break
            seen |= set(nxt_pop)
            fewest, most = min(fewest, len(nxt_pop)), max(most, len(nxt_pop))
            row += 1
        runs.append(
            Run(
                first_row=start + 1,
                last_row=row + 1,
                state=state,
                columns=sorted(seen, key=lambda c: (len(c), c)),
                fewest=fewest,
                most=most,
                deviations=deviations,
                first_value=labels[start] if labels is not None and start < len(labels) else "",
            )
        )
        row += 1
    return ConformanceMap(
        reference={get_column_letter(col + 1): dtype for col, dtype in sorted(reference.items())}, runs=runs
    )
