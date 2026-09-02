from enum import StrEnum

from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel, Field

from citadel.tabular.sheet_dump import SheetDump
from citadel.tabular.sheet_tools import AGGREGATE, RANGE


class RowRole(StrEnum):
    HEADER = "header"
    BODY = "body"
    TOTALS = "totals"
    BAND_LABEL = "band_label"
    BLANK = "blank"


class MetadataKind(StrEnum):
    TITLE = "title"
    CAPTION = "caption"
    PROVENANCE = "provenance"
    NOTE = "note"
    KEY_VALUE = "key_value"
    LEGEND = "legend"
    PROSE = "prose"


class Orientation(StrEnum):
    ROW_RECORDS = "row_records"
    CROSS_TAB = "cross_tab"
    REPEATED_GROUPS = "repeated_groups"
    MATRIX = "matrix"


class Confidence(StrEnum):
    CORROBORATED = "corroborated"
    INFERRED = "inferred"
    UNCERTAIN = "uncertain"


class RowSpan(BaseModel):
    from_row: int
    to_row: int
    role: RowRole


class Band(BaseModel):
    label: str
    from_row: int
    to_row: int


class ColumnDef(BaseModel):
    letter: str
    name: str
    index: int = 0
    header_parts: list[str] = Field(default_factory=list)
    group: str | None = None


class Region(BaseModel):
    region_id: str
    first_row: int
    last_row: int
    first_col: int
    last_col: int
    row_spans: list[RowSpan] = Field(default_factory=list)
    columns: list[ColumnDef] = Field(default_factory=list)
    key_column: str | None = None
    orientation: Orientation = Orientation.ROW_RECORDS
    group_name: str | None = None
    band_name: str | None = None
    band_column: str | None = None
    bands: list[Band] = Field(default_factory=list)
    continues_before: bool = False
    continues_after: bool = False
    confidence: Confidence = Confidence.INFERRED
    support: list[str] = Field(default_factory=list)


class MetadataItem(BaseModel):
    kind: MetadataKind
    first_row: int
    last_row: int
    first_col: int
    last_col: int
    summary: str
    region_ids: list[str] = Field(default_factory=list)


class SheetStructure(BaseModel):
    workbook: str = ""
    sheet: str = ""
    regions: list[Region] = Field(default_factory=list)
    metadata: list[MetadataItem] = Field(default_factory=list)
    unresolved: list[str] = Field(default_factory=list)
    rounds: int = 0


PREVIEW = 20
RANGE_PAD = 1
MIN_TABLE_COLUMNS = 2


class DraftSpan(BaseModel):
    rows: tuple[int, int]
    role: RowRole


class DraftRegion(BaseModel):
    id: str
    rows: tuple[int, int]
    cols: tuple[int, int]
    spans: list[DraftSpan] = Field(default_factory=list)


class DraftBlock(BaseModel):
    rows: tuple[int, int]
    cols: tuple[int, int]
    kind: MetadataKind


class Draft(BaseModel):
    regions: list[DraftRegion] = Field(default_factory=list)
    blocks: list[DraftBlock] = Field(default_factory=list)


def draft_boxes(draft: Draft) -> list[tuple[int, int, int, int]]:
    regions = [(r.rows[0], r.cols[0], r.rows[1], r.cols[1]) for r in draft.regions]
    blocks = [(b.rows[0], b.cols[0], b.rows[1], b.cols[1]) for b in draft.blocks]
    return regions + blocks


def uncovered_cells(dump: SheetDump, draft: Draft) -> list[str]:
    boxes = draft_boxes(draft)
    missed = [
        f"{get_column_letter(col)}{row}"
        for row, columns in dump.cells.items()
        for col in columns
        if not any(box[0] <= row <= box[2] and box[1] <= col <= box[3] for box in boxes)
    ]
    return missed


def review_draft(dump: SheetDump, draft: Draft) -> list[str]:
    findings: list[str] = []
    for region in draft.regions:
        seen: set[int] = set()
        doubled: set[int] = set()
        for span in region.spans:
            for row in range(min(span.rows), max(span.rows) + 1):
                if row in seen:
                    doubled.add(row)
                seen.add(row)
        missing = [row for row in range(region.rows[0], region.rows[1] + 1) if row not in seen]
        if missing:
            findings.append(f"region {region.id} leaves rows {', '.join(spans(missing)[:PREVIEW])} without a role")
        if doubled:
            findings.append(f"region {region.id} gives rows {', '.join(spans(sorted(doubled))[:PREVIEW])} two roles")
        outside = sorted(row for row in seen if not region.rows[0] <= row <= region.rows[1])
        if outside:
            findings.append(f"region {region.id} gives roles to rows outside itself: {spans(outside)[:PREVIEW]}")
    boxes = {region.id: (region.rows[0], region.cols[0], region.rows[1], region.cols[1]) for region in draft.regions}
    names = sorted(boxes)
    findings.extend(
        f"regions {a} and {b} overlap"
        for i, a in enumerate(names)
        for b in names[i + 1 :]
        if boxes_overlap(boxes[a], boxes[b])
    )
    missed = uncovered_cells(dump, draft)
    if missed:
        preview = ", ".join(missed[:PREVIEW])
        findings.append(
            f"{len(missed)} populated cells belong to no region and no block: {preview}"
            f"{' …' if len(missed) > PREVIEW else ''}"
        )
    return findings


def spans(rows: list[int]) -> list[str]:
    out: list[list[int]] = []
    for row in sorted(rows):
        if out and row == out[-1][1] + 1:
            out[-1][1] = row
        else:
            out.append([row, row])
    return [str(a) if a == b else f"{a}-{b}" for a, b in out]


def span_rows(span: RowSpan) -> range:
    return range(min(span.from_row, span.to_row), max(span.from_row, span.to_row) + 1)


def rows_by_role(region: Region, role: RowRole) -> list[int]:
    return [row for span in region.row_spans if span.role == role for row in span_rows(span)]


def body_row_numbers(region: Region) -> list[int]:
    return rows_by_role(region, RowRole.BODY)


def region_box(region: Region) -> tuple[int, int, int, int]:
    return (region.first_row, region.first_col, region.last_row, region.last_col)


def boxes_overlap(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    return a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]


def letter_index(letter: str | None) -> int:
    letters = "".join(ch for ch in (letter or "").upper() if ch.isalpha())
    return column_index_from_string(letters) if letters else 0


def coverage(dump: SheetDump, structure: SheetStructure) -> float:
    total = dump.last_row - dump.first_row + 1
    if total <= 0:
        return 1.0
    covered = {
        row
        for region in structure.regions
        for row in range(region.first_row, region.last_row + 1)
        if dump.first_row <= row <= dump.last_row
    }
    covered |= {
        row
        for item in structure.metadata
        for row in range(item.first_row, item.last_row + 1)
        if dump.first_row <= row <= dump.last_row
    }
    return len(covered) / total


def populated_columns(dump: SheetDump, region: Region) -> list[int]:
    return [
        col
        for col in range(region.first_col, region.last_col + 1)
        if any(dump.cell(row, col) for row in range(region.first_row, region.last_row + 1))
    ]


def blank_columns_inside(dump: SheetDump, region: Region) -> list[int]:
    populated = populated_columns(dump, region)
    if len(populated) < MIN_TABLE_COLUMNS:
        return []
    runs: list[list[int]] = []
    for col in range(populated[0] + 1, populated[-1]):
        if col in populated:
            continue
        if runs and runs[-1][-1] == col - 1:
            runs[-1].append(col)
        else:
            runs.append([col])
    return [col for run in runs for col in run if len(run) == 1]


def review_spans(region: Region) -> list[str]:
    seen: set[int] = set()
    doubled: set[int] = set()
    for span in region.row_spans:
        if span.to_row < span.from_row:
            return [f"region {region.region_id} has an inverted row span {span.from_row}-{span.to_row}"]
        for row in span_rows(span):
            if row in seen:
                doubled.add(row)
            seen.add(row)
    findings: list[str] = []
    missing = [row for row in range(region.first_row, region.last_row + 1) if row not in seen]
    if missing:
        listed = ", ".join(spans(missing)[:PREVIEW])
        findings.append(f"region {region.region_id} leaves rows {listed} without a role")
    outside = sorted(row for row in seen if not region.first_row <= row <= region.last_row)
    if outside:
        findings.append(f"region {region.region_id} gives roles to rows outside itself: {spans(outside)[:PREVIEW]}")
    if doubled:
        listed = ", ".join(spans(sorted(doubled))[:PREVIEW])
        findings.append(f"region {region.region_id} gives rows {listed} more than one role")
    if not body_row_numbers(region):
        findings.append(f"region {region.region_id} has no rows holding records")
    return findings


def review_shape(dump: SheetDump, region: Region) -> list[str]:
    findings: list[str] = []
    if region.first_row < dump.first_row or region.last_row > dump.last_row:
        findings.append(f"region {region.region_id} covers rows outside the worksheet shown")
    populated = populated_columns(dump, region)
    if len(populated) < MIN_TABLE_COLUMNS:
        findings.append(
            f"region {region.region_id} occupies {len(populated)} populated column(s); "
            "a table needs at least two, and separators inside a cell's text do not make columns"
        )
    return findings


def advise_shape(dump: SheetDump, region: Region) -> list[str]:
    notes: list[str] = []
    blanks = blank_columns_inside(dump, region)
    if blanks:
        listed = ", ".join(get_column_letter(col) for col in blanks)
        notes.append(
            f"region {region.region_id} spans blank column(s) {listed}; keep it as one region if a header "
            "names the columns on both sides, otherwise these are separate regions"
        )
    if region.key_column is None:
        notes.append(f"region {region.region_id} names no key column; say which column identifies its rows")
    elif not region.first_col <= letter_index(region.key_column) <= region.last_col:
        notes.append(f"region {region.region_id} has a key column {region.key_column} outside its own columns")
    if region.band_column and not region.first_col <= letter_index(region.band_column) <= region.last_col:
        notes.append(f"region {region.region_id} has a band column {region.band_column} outside its own columns")
    return notes


def advise_structure(dump: SheetDump, structure: SheetStructure) -> list[str]:
    return [note for region in structure.regions for note in advise_shape(dump, region)]


def columns_overlap(a: Region, b: Region) -> bool:
    return a.first_col <= b.last_col and b.first_col <= a.last_col


def review_between(structure: SheetStructure) -> list[str]:
    findings = [
        f"regions {a.region_id} and {b.region_id} overlap"
        for i, a in enumerate(structure.regions)
        for j, b in enumerate(structure.regions)
        if i < j and boxes_overlap(region_box(a), region_box(b))
    ]
    known = {region.region_id for region in structure.regions}
    for item in structure.metadata:
        unknown = [name for name in item.region_ids if name not in known]
        if unknown:
            findings.append(f"metadata at rows {item.first_row}-{item.last_row} names unknown regions {unknown}")
    for region in structure.regions:
        headers = set(rows_by_role(region, RowRole.HEADER))
        for other in structure.regions:
            if other.region_id == region.region_id or not columns_overlap(region, other):
                continue
            inside = sorted(headers & set(body_row_numbers(other)))
            if inside:
                findings.append(
                    f"region {region.region_id} marks rows {spans(inside)[:PREVIEW]} as headers while "
                    f"region {other.region_id} treats them as records; stacked tables each keep their own header"
                )
    return findings


def review_structure(dump: SheetDump, structure: SheetStructure) -> list[str]:
    findings: list[str] = []
    for region in structure.regions:
        findings.extend(review_spans(region))
        findings.extend(review_shape(dump, region))
    findings.extend(review_between(structure))
    return findings


def normalise_ref(ref: str) -> str:
    return ref.replace("$", "").replace("'", "").upper().split("!")[-1].strip()


def extent_box(extent: str) -> tuple[int, int, int, int] | None:
    match = RANGE.match(extent.strip().upper().replace(" ", ""))
    if not match or not match.group(4):
        return None
    col_a = column_index_from_string(match.group(1))
    col_b = column_index_from_string(match.group(3))
    return (
        min(int(match.group(2)), int(match.group(4))),
        min(col_a, col_b),
        max(int(match.group(2)), int(match.group(4))),
        max(col_a, col_b),
    )


def declared_ranges(dump: SheetDump) -> dict[str, str]:
    out = {normalise_ref(value): f"defined name {name}" for name, value in dump.meta.defined_names.items()}
    for merged in dump.meta.merged:
        out.setdefault(normalise_ref(merged), f"merged range {merged}")
    return out


def aggregate_ranges(dump: SheetDump) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in dump.cells.values():
        for cell in row.values():
            if cell.formula is None:
                continue
            for match in AGGREGATE.finditer(cell.formula):
                out[normalise_ref(match.group(2))] = f"{match.group(1)} at {cell.ref}"
    return out


def corroboration(dump: SheetDump, region: Region) -> tuple[Confidence, list[str]]:
    box = region_box(region)
    found: list[str] = []
    for ref, source in declared_ranges(dump).items():
        other = extent_box(ref)
        if other is not None and boxes_overlap(box, other) and (other[0], other[2]) == (box[0], box[2]):
            found.append(f"{source} matches rows {box[0]}-{box[2]}")
    rows = body_row_numbers(region)
    if rows:
        low, high = min(rows), max(rows)
        for ref, source in aggregate_ranges(dump).items():
            other = extent_box(ref)
            if other is None or not box[1] <= other[1] <= box[3]:
                continue
            if other[0] <= low <= other[0] + RANGE_PAD and other[2] - RANGE_PAD <= high <= other[2]:
                found.append(f"{source} covers rows {low}-{high}")
    if found:
        return Confidence.CORROBORATED, found[:4]
    return Confidence.INFERRED, []


def score_structure(dump: SheetDump, structure: SheetStructure) -> dict[str, int]:
    tally: dict[str, int] = {}
    for region in structure.regions:
        level, support = corroboration(dump, region)
        region.confidence = level
        region.support = support
        tally[level.value] = tally.get(level.value, 0) + 1
    return tally


COMPUTED_FIELDS = {
    "SheetStructure": ("workbook", "sheet", "rounds"),
    "Region": ("confidence", "support", "continues_before", "continues_after"),
    "ColumnDef": ("index",),
}

REQUIRED_FIELDS = {
    "SheetStructure": ("regions", "metadata"),
    "Region": (
        "region_id",
        "first_row",
        "last_row",
        "first_col",
        "last_col",
        "row_spans",
        "columns",
        "key_column",
        "orientation",
        "group_name",
        "band_name",
        "band_column",
        "bands",
    ),
    "MetadataItem": ("kind", "first_row", "last_row", "first_col", "last_col", "summary", "region_ids"),
    "ColumnDef": ("letter", "name"),
}

MIN_ITEMS = {"Region": {"row_spans": 1, "columns": 1}}


def prune(node: dict, names: tuple[str, ...]) -> None:
    for name in names:
        node.get("properties", {}).pop(name, None)
        if name in node.get("required", []):
            node["required"].remove(name)


def require(node: dict, names: tuple[str, ...]) -> None:
    present = node.get("properties", {})
    required = node.setdefault("required", [])
    for name in names:
        if name in present and name not in required:
            required.append(name)


def set_min_items(node: dict, limits: dict[str, int]) -> None:
    properties = node.get("properties", {})
    for name, minimum in limits.items():
        field = properties.get(name)
        if isinstance(field, dict) and field.get("type") == "array":
            field["minItems"] = minimum


def request_schema() -> dict:
    schema = SheetStructure.model_json_schema()
    prune(schema, COMPUTED_FIELDS["SheetStructure"])
    require(schema, REQUIRED_FIELDS["SheetStructure"])
    defs = schema.get("$defs", {})
    for name, fields in COMPUTED_FIELDS.items():
        if (target := defs.get(name)) is not None:
            prune(target, fields)
    for name, fields in REQUIRED_FIELDS.items():
        if (target := defs.get(name)) is not None:
            require(target, fields)
    for name, limits in MIN_ITEMS.items():
        if (target := defs.get(name)) is not None:
            set_min_items(target, limits)
    return schema


def fill_edges(dump: SheetDump, region: Region) -> None:
    if not dump.windowed:
        region.continues_before = False
        region.continues_after = False
        return
    if region.first_row <= dump.first_row and dump.first_row > dump.sheet_first_row:
        region.continues_before = True
    if region.last_row >= dump.last_row and dump.last_row < dump.sheet_last_row:
        region.continues_after = True


def fill_computed(dump: SheetDump, structure: SheetStructure) -> SheetStructure:
    structure.workbook = dump.workbook
    structure.sheet = dump.sheet
    for region in structure.regions:
        fill_edges(dump, region)
        for column in region.columns:
            column.index = letter_index(column.letter)
    return structure
