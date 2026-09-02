import re
from collections import Counter

from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel, Field

from citadel.tabular.sheet_dump import Cell, SheetDump

RANGE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?$")
AGGREGATE = re.compile(r"\b(SUM|AVERAGE|AVERAGEA|COUNT|COUNTA|SUBTOTAL|MIN|MAX)\(([^()]*?:[^()]*?)\)")
MAX_HITS = 200
MAX_VALUES = 50


class FindHit(BaseModel):
    row: int
    text: str


class FindResult(BaseModel):
    pattern: str
    count: int
    truncated: bool
    rows: list[FindHit] = Field(default_factory=list)


class ValueCount(BaseModel):
    value: str
    count: int


class ColumnValues(BaseModel):
    column: str
    populated: int
    distinct: int
    truncated: bool
    values: list[ValueCount] = Field(default_factory=list)


class RangeResult(BaseModel):
    ref: str
    rows: int
    columns: int
    populated: int
    cells: list[Cell] = Field(default_factory=list)


class RowCount(BaseModel):
    row: int
    cells: int


class RowOccupancy(BaseModel):
    from_row: int
    to_row: int
    columns: str
    rows: list[RowCount] = Field(default_factory=list)


class ColumnCount(BaseModel):
    column: str
    cells: int
    blank: bool
    kinds: str


class ColumnOccupancy(BaseModel):
    from_row: int
    to_row: int
    columns: list[ColumnCount] = Field(default_factory=list)


class ToolError(Exception):
    pass


TOOLS = {
    "find": ("pattern", "from_row", "to_row"),
    "column_values": ("column", "from_row", "to_row"),
    "read_range": ("ref",),
    "row_occupancy": ("from_row", "to_row", "first_col", "last_col"),
    "column_occupancy": ("from_row", "to_row", "first_col", "last_col"),
}


def dispatch(dump: SheetDump, name: str, arguments: dict) -> BaseModel:
    if name not in TOOLS:
        message = f"there is no tool called {name}; available: {', '.join(sorted(TOOLS))}"
        raise ToolError(message)
    allowed = TOOLS[name]
    unknown = sorted(set(arguments) - set(allowed))
    if unknown:
        message = f"{name} takes {', '.join(allowed)}; it has no argument {', '.join(unknown)}"
        raise ToolError(message)
    return globals()[name](dump, **arguments)


def render_result(result: BaseModel) -> str:
    if isinstance(result, RowOccupancy):
        counts = "  ".join(f"{row.row}:{row.cells}" for row in result.rows)
        return f"rows {result.from_row}-{result.to_row} over columns {result.columns}, cells per row:\n{counts}"
    if isinstance(result, ColumnOccupancy):
        counts = "  ".join(
            f"{col.column}:{'blank' if col.blank else col.cells}" for col in result.columns
        )
        return f"rows {result.from_row}-{result.to_row}, cells per column:\n{counts}"
    if isinstance(result, ColumnValues):
        values = "  ".join(f"{item.value}×{item.count}" for item in result.values)
        return (
            f"column {result.column}: {result.populated} populated, {result.distinct} distinct"
            f"{' (truncated)' if result.truncated else ''}\n{values}"
        )
    if isinstance(result, FindResult):
        hits = "\n".join(f"{hit.row}|{hit.text}" for hit in result.rows)
        return f"{result.count} rows match {result.pattern!r}{' (truncated)' if result.truncated else ''}\n{hits}"
    if isinstance(result, RangeResult):
        cells = "  ".join(f"{cell.ref}={cell.value}" for cell in result.cells)
        return f"{result.ref}: {result.populated} populated of {result.rows}x{result.columns}\n{cells}"
    return result.model_dump_json()


def bounds(dump: SheetDump, from_row: int | None, to_row: int | None) -> tuple[int, int]:
    start = dump.first_row if from_row is None else max(from_row, dump.first_row)
    end = dump.last_row if to_row is None else min(to_row, dump.last_row)
    return start, end


def find(dump: SheetDump, pattern: str, from_row: int | None = None, to_row: int | None = None) -> FindResult:
    try:
        matcher = re.compile(pattern)
    except re.error as exc:
        message = f"invalid pattern: {exc}"
        raise ToolError(message) from exc
    start, end = bounds(dump, from_row, to_row)
    hits = [
        FindHit(row=row, text=dump.lines[row])
        for row in range(start, end + 1)
        if row in dump.lines and matcher.search(dump.lines[row])
    ]
    return FindResult(pattern=pattern, count=len(hits), truncated=len(hits) > MAX_HITS, rows=hits[:MAX_HITS])


def column_values(
    dump: SheetDump,
    column: str | int,
    from_row: int | None = None,
    to_row: int | None = None,
    limit: int = MAX_VALUES,
) -> ColumnValues:
    letters = get_column_letter(column) if isinstance(column, int) else str(column).strip().upper()
    if not letters.isalpha():
        message = f"invalid column: {column}"
        raise ToolError(message)
    index = column_index_from_string(letters)
    start, end = bounds(dump, from_row, to_row)
    counter: Counter[str] = Counter()
    for row in range(start, end + 1):
        cell = dump.cell(row, index)
        if cell is not None and cell.value:
            counter[cell.value] += 1
    ordered = counter.most_common()
    return ColumnValues(
        column=letters,
        populated=sum(counter.values()),
        distinct=len(counter),
        truncated=len(ordered) > limit,
        values=[ValueCount(value=value, count=count) for value, count in ordered[:limit]],
    )


def column_bounds(dump: SheetDump, first_col: int | None, last_col: int | None) -> tuple[int, int]:
    start = dump.first_col if first_col is None else max(first_col, dump.first_col)
    end = dump.last_col if last_col is None else min(last_col, dump.last_col)
    return start, end


def row_occupancy(
    dump: SheetDump,
    from_row: int | None = None,
    to_row: int | None = None,
    first_col: int | None = None,
    last_col: int | None = None,
) -> RowOccupancy:
    start, end = bounds(dump, from_row, to_row)
    col_start, col_end = column_bounds(dump, first_col, last_col)
    counts = [
        RowCount(row=row, cells=sum(1 for col in range(col_start, col_end + 1) if dump.cell(row, col) is not None))
        for row in range(start, end + 1)
    ]
    return RowOccupancy(
        from_row=start,
        to_row=end,
        columns=f"{get_column_letter(col_start)}-{get_column_letter(col_end)}",
        rows=counts,
    )


def column_occupancy(
    dump: SheetDump,
    from_row: int | None = None,
    to_row: int | None = None,
    first_col: int | None = None,
    last_col: int | None = None,
) -> ColumnOccupancy:
    start, end = bounds(dump, from_row, to_row)
    col_start, col_end = column_bounds(dump, first_col, last_col)
    columns: list[ColumnCount] = []
    for col in range(col_start, col_end + 1):
        cells = [cell for row in range(start, end + 1) if (cell := dump.cell(row, col)) is not None]
        kinds = Counter("formula" if cell.formula else "text" for cell in cells)
        columns.append(
            ColumnCount(
                column=get_column_letter(col),
                cells=len(cells),
                blank=not cells,
                kinds="/".join(f"{name}:{n}" for name, n in kinds.most_common()) or "-",
            )
        )
    return ColumnOccupancy(from_row=start, to_row=end, columns=columns)


def read_range(dump: SheetDump, ref: str) -> RangeResult:
    match = RANGE.match(ref.strip().upper().replace(" ", ""))
    if not match:
        message = f"invalid range: {ref}"
        raise ToolError(message)
    col_a = column_index_from_string(match.group(1))
    row_a = int(match.group(2))
    col_b = column_index_from_string(match.group(3)) if match.group(3) else col_a
    row_b = int(match.group(4)) if match.group(4) else row_a
    row_start, row_end = min(row_a, row_b), max(row_a, row_b)
    col_start, col_end = min(col_a, col_b), max(col_a, col_b)
    cells = [
        cell
        for row in range(row_start, row_end + 1)
        for column in range(col_start, col_end + 1)
        if (cell := dump.cell(row, column)) is not None
    ]
    normalised = f"{get_column_letter(col_start)}{row_start}:{get_column_letter(col_end)}{row_end}"
    return RangeResult(
        ref=normalised,
        rows=row_end - row_start + 1,
        columns=col_end - col_start + 1,
        populated=len(cells),
        cells=cells,
    )
