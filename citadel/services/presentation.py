from html import escape
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from citadel.schemas.content import Block
from citadel.services.tabular import grid_from_html

_DRAWINGML_TEXT = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
_EMU_PER_PX = 9525
_MIN_PICTURE_PX = 32


def _keep_picture(shape: Any) -> bool:
    width, height = getattr(shape, "width", None), getattr(shape, "height", None)
    if width is None or height is None:
        return True
    return not (width < _MIN_PICTURE_PX * _EMU_PER_PX and height < _MIN_PICTURE_PX * _EMU_PER_PX)


def _bbox(shape: Any) -> list[float] | None:
    left, top = getattr(shape, "left", None), getattr(shape, "top", None)
    if left is None or top is None:
        return None
    width = getattr(shape, "width", None) or 0
    height = getattr(shape, "height", None) or 0
    return [float(left), float(top), float(left + width), float(top + height)]


def _ordered(shapes: Any) -> list:
    return sorted(shapes, key=lambda shape: (shape.top or 0, shape.left or 0))


def _cell_tag(row_index: int, first_row_is_header: bool) -> str:
    return "th" if row_index == 0 and first_row_is_header else "td"


def _table_html(table: Any) -> str:
    rows: list[str] = []
    for row_index, row in enumerate(table.rows):
        tag = _cell_tag(row_index, bool(table.first_row))
        cells = "".join(f"<{tag}>{escape(cell.text.strip())}</{tag}>" for cell in row.cells)
        rows.append(f"<tr>{cells}</tr>")
    return "<table>" + "".join(rows) + "</table>"


def _series_cell(values: list, index: int) -> str:
    if index >= len(values) or values[index] is None:
        return "<td></td>"
    return f"<td>{escape(str(values[index]))}</td>"


def _chart_html(chart: Any) -> str:
    categories = [str(category) for category in chart.plots[0].categories] if chart.plots else []
    series = list(chart.series)
    header = "".join(f"<th>{escape(str(item.name or ''))}</th>" for item in series)
    values = [list(item.values) for item in series]
    rows = [f"<tr><th>category</th>{header}</tr>"]
    for index, category in enumerate(categories):
        cells = "".join(_series_cell(column, index) for column in values)
        rows.append(f"<tr><td>{escape(category)}</td>{cells}</tr>")
    return "<table>" + "".join(rows) + "</table>"


def _alt_text(shape: Any) -> str:
    for element in shape._element.iter():
        if element.tag.endswith("}cNvPr"):
            return (element.get("descr") or "").strip()
    return ""


def _graphic_text(shape: Any) -> str:
    parts = [node.text.strip() for node in shape._element.iter(_DRAWINGML_TEXT) if node.text and node.text.strip()]
    return " ".join(parts)


_Item = tuple[Block, bytes | None]


def _text_frame_blocks(shape: Any, slide_index: int, is_title: bool) -> list[_Item]:
    bbox = _bbox(shape)
    if is_title:
        text = shape.text_frame.text.strip()
        return [(Block(page_idx=slide_index, type="title", text=text, text_level=1, bbox=bbox), None)] if text else []
    items: list[_Item] = []
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text.strip()
        if text:
            items.append((Block(page_idx=slide_index, type="text", text=text, bbox=bbox), None))
    return items


def _shape_items(shape: Any, slide_index: int, title_id: int | None) -> list[_Item]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [item for member in _ordered(shape.shapes) for item in _shape_items(member, slide_index, title_id)]
    if shape.has_table:
        html = _table_html(shape.table)
        block = Block(page_idx=slide_index, type="table", text=html, bbox=_bbox(shape), grid=grid_from_html(html))
        return [(block, None)]
    if shape.has_chart:
        html = _chart_html(shape.chart)
        block = Block(page_idx=slide_index, type="table", text=html, bbox=_bbox(shape), grid=grid_from_html(html))
        return [(block, None)]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        if not _keep_picture(shape):
            return []
        block = Block(page_idx=slide_index, type="image", text=_alt_text(shape), bbox=_bbox(shape))
        return [(block, shape.image.blob)]
    if shape.has_text_frame:
        return _text_frame_blocks(shape, slide_index, shape.shape_id == title_id)
    graphic = _graphic_text(shape)
    return [(Block(page_idx=slide_index, type="text", text=graphic, bbox=_bbox(shape)), None)] if graphic else []


def _notes_items(slide: Any, slide_index: int) -> list[_Item]:
    if not slide.has_notes_slide:
        return []
    notes = slide.notes_slide.notes_text_frame.text.strip()
    return [(Block(page_idx=slide_index, type="text", text=notes), None)] if notes else []


def _slide_items(slide: Any, slide_index: int) -> list[_Item]:
    title = slide.shapes.title
    title_id = title.shape_id if title is not None else None
    items = [item for shape in _ordered(slide.shapes) for item in _shape_items(shape, slide_index, title_id)]
    return items + _notes_items(slide, slide_index)


def parse_pptx(data: bytes) -> list[tuple[list[Block], dict[int, bytes]]]:
    presentation = Presentation(BytesIO(data))
    slides: list[tuple[list[Block], dict[int, bytes]]] = []
    for index, slide in enumerate(presentation.slides):
        items = _slide_items(slide, index)
        blocks = [block for block, _ in items]
        images = {i: data for i, (_, data) in enumerate(items) if data is not None}
        slides.append((blocks, images))
    return slides
