from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.cell.cell import Cell as OpenpyxlCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from citadel.tabular.materialize import MaterializedTable

type RawCellValue = str | int | float | bool | datetime | None


@dataclass
class Cell:
    row: int
    col: int
    value: RawCellValue
    number_format: str
    bold: bool
    filled: bool
    bordered: bool
    formula: str | None = None
    comment: str | None = None


def cell_value(cell: Cell) -> RawCellValue:
    return cell.value if cell.value is not None else cell.formula


@dataclass
class SheetTable:
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    header_row_count: int


@dataclass
class SheetMetadata:
    defined_names: dict[str, str] = field(default_factory=dict)
    hidden_rows: list[int] = field(default_factory=list)
    hidden_cols: list[str] = field(default_factory=list)
    freeze: str | None = None
    autofilter: str | None = None
    validation: list[str] = field(default_factory=list)
    cond_format: list[str] = field(default_factory=list)
    hyperlinks: dict[str, str] = field(default_factory=dict)


@dataclass
class SheetExtraction:
    sheet_no: int
    sheet_name: str
    max_row: int
    max_col: int
    cells: list[Cell]
    tables: list[SheetTable]
    images: list[bytes]
    merges: list[tuple[int, int, int, int]]
    state: str = "visible"
    metadata: SheetMetadata = field(default_factory=SheetMetadata)


@dataclass
class Region:
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    cells: list[Cell]


@dataclass
class SheetText:
    sheet_no: int
    text: str


type SheetItem = MaterializedTable | SheetText


def _style_flags(cell: OpenpyxlCell) -> tuple[bool, bool, bool]:
    bold = bool(cell.font and cell.font.bold)
    fill = cell.fill
    filled = bool(fill and fill.patternType and fill.patternType != "none")
    border = cell.border
    bordered = bool(
        border and any(side and side.style for side in (border.left, border.right, border.top, border.bottom))
    )
    return bold, filled, bordered


def _cell_comment(cell: OpenpyxlCell) -> str | None:
    if cell.comment is None:
        return None
    return cell.comment.text.strip() or None


def _cell_formula(cell: OpenpyxlCell) -> str | None:
    value = cell.value
    return value if isinstance(value, str) and value.startswith("=") else None


def _capture_cells(values_sheet: Worksheet, formulas_sheet: Worksheet) -> list[Cell]:
    cells: list[Cell] = []
    for row in formulas_sheet.iter_rows():
        for cell in row:
            formula = _cell_formula(cell)
            comment = _cell_comment(cell)
            value = values_sheet.cell(row=cell.row, column=cell.column).value
            if value is None and formula is None and comment is None:
                continue
            bold, filled, bordered = _style_flags(cell)
            cells.append(
                Cell(
                    row=cell.row,
                    col=cell.column,
                    value=value,
                    number_format=cell.number_format or "General",
                    bold=bold,
                    filled=filled,
                    bordered=bordered,
                    formula=formula,
                    comment=comment,
                )
            )
    return cells


def _capture_tables(worksheet: Worksheet) -> list[SheetTable]:
    tables: list[SheetTable] = []
    for table in worksheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        tables.append(
            SheetTable(
                min_row=min_row,
                min_col=min_col,
                max_row=max_row,
                max_col=max_col,
                header_row_count=table.headerRowCount or 1,
            )
        )
    return tables


def _capture_images(worksheet: Worksheet) -> list[bytes]:
    return [image._data() for image in worksheet._images]


def _capture_merges(worksheet: Worksheet) -> list[tuple[int, int, int, int]]:
    return sorted(
        (cells.min_row, cells.min_col, cells.max_row, cells.max_col) for cells in worksheet.merged_cells.ranges
    )


def _capture_metadata(worksheet: Worksheet, defined_names: dict[str, str]) -> SheetMetadata:
    rows = [index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden]
    cols = [name for name, dimension in worksheet.column_dimensions.items() if dimension.hidden]
    return SheetMetadata(
        defined_names=defined_names,
        hidden_rows=sorted(rows),
        hidden_cols=sorted(cols, key=column_index_from_string),
        freeze=worksheet.freeze_panes,
        autofilter=worksheet.auto_filter.ref,
        validation=[str(rule.sqref) for rule in worksheet.data_validations.dataValidation],
        cond_format=[str(rule.sqref) for rule in worksheet.conditional_formatting],
        hyperlinks={
            cell.coordinate: str(cell.hyperlink.target)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.hyperlink is not None
        },
    )


def _sheet_defined_names(workbook: Workbook, worksheet: Worksheet) -> dict[str, str]:
    quoted = f"'{worksheet.title}'!"
    plain = f"{worksheet.title}!"
    found: dict[str, str] = {}
    for scope, items in [("", workbook.defined_names.items())] + [
        (f"[{sheet.title}]", sheet.defined_names.items()) for sheet in workbook.worksheets
    ]:
        for name, entry in items:
            value = str(entry.value)
            if value.startswith((quoted, plain)) or scope == f"[{worksheet.title}]":
                found[f"{scope}{name}"] = value
    return found


def extract_sheet(
    values_sheet: Worksheet,
    formulas_sheet: Worksheet,
    sheet_no: int,
    metadata: SheetMetadata | None = None,
) -> SheetExtraction:
    return SheetExtraction(
        state=formulas_sheet.sheet_state,
        metadata=metadata or SheetMetadata(),
        sheet_no=sheet_no,
        sheet_name=formulas_sheet.title,
        max_row=formulas_sheet.max_row or 0,
        max_col=formulas_sheet.max_column or 0,
        cells=_capture_cells(values_sheet, formulas_sheet),
        tables=_capture_tables(formulas_sheet),
        images=_capture_images(formulas_sheet),
        merges=_capture_merges(formulas_sheet),
    )


def sheet_names(data: bytes) -> list[str]:
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def load_all_sheets(data: bytes) -> list[SheetExtraction]:
    values_workbook = openpyxl.load_workbook(BytesIO(data), data_only=True)
    formulas_workbook = openpyxl.load_workbook(BytesIO(data), data_only=False)
    try:
        return [
            extract_sheet(
                values_workbook.worksheets[index],
                formulas_sheet,
                index + 1,
                _capture_metadata(formulas_sheet, _sheet_defined_names(formulas_workbook, formulas_sheet)),
            )
            for index, formulas_sheet in enumerate(formulas_workbook.worksheets)
        ]
    finally:
        values_workbook.close()
        formulas_workbook.close()


DUMP_VERSION = 1
FENCE = "````"
FORMULA_OPEN = "\u2039"
FORMULA_CLOSE = "\u203a"


def _dump_text(value: str) -> str:
    if not value.strip():
        return ""
    escaped = value.replace("\r\n", "\\n").replace("\n", "\\n").replace("\r", "\\n")
    return escaped.replace("\t", " ").replace("|", "\\|")


def _dump_value(value: RawCellValue) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        empty_time = (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0)
        return value.date().isoformat() if empty_time else value.isoformat(sep=" ")
    if isinstance(value, float):
        return f"{value:.10g}"
    if isinstance(value, int):
        return str(value)
    return _dump_text(str(value))


def _dump_cell(cell: Cell) -> str:
    rendered = _dump_value(cell.value)
    if cell.formula is None:
        return rendered
    formula = cell.formula.replace("\n", " ").replace("|", "\\|")
    marked = f"{FORMULA_OPEN}{formula}{FORMULA_CLOSE}"
    return f"{rendered} {marked}" if rendered else marked


def _column_runs(items: list[tuple[int, str]]) -> list[tuple[int, int, str]]:
    runs: list[list] = []
    for index, key in items:
        if runs and runs[-1][2] == key and index == runs[-1][1] + 1:
            runs[-1][1] = index
        else:
            runs.append([index, index, key])
    return [(a, b, key) for a, b, key in runs]


def _span_label(first: int, last: int) -> str:
    return get_column_letter(first) if first == last else f"{get_column_letter(first)}-{get_column_letter(last)}"


def _kind(cell: Cell) -> str:
    if cell.formula is not None:
        return "formula"
    value = cell.value
    if isinstance(value, datetime):
        return "date"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    return "text"


def render_sheet_dump(sheet: SheetExtraction, workbook: str) -> str:
    grid = {(cell.row, cell.col): cell for cell in sheet.cells if _dump_cell(cell)}
    lines = [f"# {workbook} / {sheet.sheet_name}", ""]
    if not grid:
        lines += ["Empty sheet: no populated cells.", ""]
        return "\n".join(lines)
    rows = [key[0] for key in grid]
    cols = [key[1] for key in grid]
    first_row, last_row, first_col, last_col = min(rows), max(rows), min(cols), max(cols)
    extent = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"
    formulas = sum(1 for cell in grid.values() if cell.formula is not None)
    lines += [
        f"- source: `{workbook}` sheet `{sheet.sheet_name}` (state: {sheet.state})",
        (
            f"- extent: `{extent}`  rows {first_row}-{last_row} ({last_row - first_row + 1})  "
            f"cols {_span_label(first_col, last_col)} ({first_col}-{last_col}, {last_col - first_col + 1})"
        ),
        f"- populated cells: {len(grid)}  formula cells: {formulas}",
        "",
        "## GRID",
        "",
        FENCE,
        "|".join(["#"] + [f"{get_column_letter(col)}({col})" for col in range(first_col, last_col + 1)]),
    ]
    for row in range(first_row, last_row + 1):
        values = [_dump_cell(grid[row, col]) if (row, col) in grid else "" for col in range(first_col, last_col + 1)]
        while values and not values[-1]:
            values.pop()
        lines.append("|".join([str(row), *values]))
    lines += [FENCE, "", "## METADATA", ""]

    lines.extend(_dump_metadata(sheet, grid, first_row, last_row, first_col, last_col))
    lines.append("")
    return "\n".join(lines)


def _column_summaries(
    grid: dict[tuple[int, int], Cell], first_row: int, last_row: int, first_col: int, last_col: int
) -> tuple[list[tuple[int, str]], list[tuple[int, str]]]:
    types: list[tuple[int, str]] = []
    formats: list[tuple[int, str]] = []
    for col in range(first_col, last_col + 1):
        column = [grid[row, col] for row in range(first_row, last_row + 1) if (row, col) in grid]
        if not column:
            continue
        counts = Counter(_kind(cell) for cell in column)
        types.append((col, "/".join(f"{name}:{n}" for name, n in counts.most_common())))
        seen = sorted({cell.number_format for cell in column} - {"General"})
        if seen:
            formats.append((col, ",".join(seen[:3])))
    return types, formats


def _dump_metadata(
    sheet: SheetExtraction,
    grid: dict[tuple[int, int], Cell],
    first_row: int,
    last_row: int,
    first_col: int,
    last_col: int,
) -> list[str]:
    lines: list[str] = []
    merged = [
        f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}"
        for r0, c0, r1, c1 in sorted(sheet.merges, key=lambda box: (box[0], -box[1]))
    ]
    lines.append(f"- MERGED: {' | '.join(merged) or '-'}")
    types, formats = _column_summaries(grid, first_row, last_row, first_col, last_col)
    spans = "  ".join(f"{_span_label(a, b)} {k}" for a, b, k in _column_runs(types))
    lines.append(f"- TYPES: {spans or '-'}")
    spans = "  ".join(f"{_span_label(a, b)} {k}" for a, b, k in _column_runs(formats))
    lines.append(f"- FORMATS: {spans or '-'}")
    meta = sheet.metadata
    lines.extend(
        [
            f"- HIDDEN: rows {meta.hidden_rows or '-'}  cols {meta.hidden_cols or '-'}",
            f"- FREEZE: {meta.freeze or '-'}",
            f"- AUTOFILTER: {meta.autofilter or '-'}",
        ]
    )
    listobjects = [
        f"{get_column_letter(item.min_col)}{item.min_row}:{get_column_letter(item.max_col)}{item.max_row}"
        for item in sheet.tables
    ]
    lines.append(f"- LISTOBJECTS: {listobjects or '-'}")
    names = " | ".join(f"{name}={value}" for name, value in meta.defined_names.items())
    lines.extend(
        [
            f"- DEFINED NAMES: {names or '-'}",
            f"- VALIDATION: {' | '.join(meta.validation) or '-'}",
            f"- CONDFMT: {' | '.join(meta.cond_format) or '-'}",
        ]
    )
    comments = " | ".join(
        f"{get_column_letter(cell.col)}{cell.row}={_dump_text(cell.comment)[:80]}"
        for cell in sheet.cells
        if cell.comment
    )
    lines.append(f"- COMMENTS: {comments or '-'}")
    links = " | ".join(f"{ref}={target}" for ref, target in meta.hyperlinks.items())
    lines.append(f"- HYPERLINKS: {links or '-'}")
    by_row: dict[int, list[int]] = {}
    for cell in sheet.cells:
        if cell.bold and (cell.row, cell.col) in grid:
            by_row.setdefault(cell.row, []).append(cell.col)
    signature = [
        (row, ",".join(get_column_letter(col) for col in sorted(cols))) for row, cols in sorted(by_row.items())
    ]
    bold_text = "  ".join(f"rows {a}-{b} cols {key}" for a, b, key in _column_runs(signature))
    lines.append(f"- BOLD: {bold_text or '-'}")
    return lines
